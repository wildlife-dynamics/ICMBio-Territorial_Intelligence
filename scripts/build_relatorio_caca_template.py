"""Build the Relatório de Inteligência Territorial — Eventos de Caça docxtpl
(Jinja2) template FROM SCRATCH via python-docx.

Unlike ICMBio-patrol_analysis's relatorio_bimestral_template.docx (converted
from a real partner-supplied mockup .docx), there is no partner-supplied
source document for this report — only the fully rendered reference PDF
(Relatorio_Inteligencia_Territorial_Caca_EarthRanger_Versao_Final.pdf). This
script therefore builds the template directly, section-by-section matching
that PDF's own numbering and structure (and the PRD's — the two agree here).
No attempt is made to reproduce the partner's exact letterhead/logo
graphics, since no source asset for those exists in this repo.

Data-driven sections (charts/maps/tables/photos) use docxtpl tags, exactly
matching the context keys produced by
ecoscope_workflows_ext_icmbio.tasks.prepare_hunting_report_context /
generate_hunting_report. Narrative sections (Objetivo, Fonte dos dados e
metodologia, Relações com acesso/limites/cursos d'água, Tendências
observadas, Avaliação Estratégica, Recomendações) are plain headings with a
blank editable paragraph — per the PRD, that text is written by the user,
not generated.

Run with:
    pixi run -e default python scripts/build_relatorio_caca_template.py
(from ICMBio-patrol_analysis/ecoscope-workflows-patrol-analysis-workflow,
or any other env with python-docx installed.)
"""

from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.table import Table, _Row

DST = Path(__file__).parent.parent / "resources" / "templates" / "relatorio_inteligencia_territorial_caca_template.docx"

TEAL = RGBColor(0x1F, 0x6F, 0x5C)
ORANGE = RGBColor(0xC8, 0x6A, 0x1F)
GREY = RGBColor(0x55, 0x55, 0x55)


# ── low-level helpers (same pattern as
#    ICMBio-patrol_analysis/scripts/build_relatorio_bimestral_template.py's
#    make_row_loop / clone_row, generalized to build from scratch rather
#    than edit an existing mockup) ──────────────────────────────────────────


def clear_cell(cell):
    cell.text = ""
    return cell.paragraphs[0]


def set_cell_text(cell, text: str, bold: bool = False):
    p = clear_cell(cell)
    r = p.add_run(text)
    r.bold = bold
    return r


def clone_row(table: Table, ref_row, after: bool) -> _Row:
    import copy

    new_tr = copy.deepcopy(ref_row._tr)
    if after:
        ref_row._tr.addnext(new_tr)
    else:
        ref_row._tr.addprevious(new_tr)
    return _Row(new_tr, table)


def add_table_with_row_loop(doc: docx.Document, headers: list[str], loop_var: str, cell_exprs: list[str]) -> Table:
    """A header row + a docxtpl row-loop (start row / data row / end row) —
    same 3-row pattern as the bimestral template's make_row_loop."""
    table = doc.add_table(rows=2, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)

    data_row = table.rows[1]
    loop_start = clone_row(table, data_row, after=False)
    loop_end = clone_row(table, data_row, after=True)

    set_cell_text(loop_start.cells[0], "{%tr for item in " + loop_var + " %}")
    for c in loop_start.cells[1:]:
        clear_cell(c)

    for cell, expr in zip(data_row.cells, cell_exprs):
        set_cell_text(cell, expr)

    set_cell_text(loop_end.cells[0], "{%tr endfor %}")
    for c in loop_end.cells[1:]:
        clear_cell(c)

    return table


def add_heading(doc: docx.Document, number: str, text: str, level: int = 1):
    heading = doc.add_heading(level=level)
    run = heading.add_run(f"{number}. {text}" if number else text)
    run.font.color.rgb = TEAL
    return heading


def add_blank_editable_paragraph(doc: docx.Document, placeholder: str):
    """A placeholder paragraph for narrative text the user adds in Word —
    per the PRD, descriptive/analytical text and recommendations are
    written by the user, not generated."""
    p = doc.add_paragraph()
    r = p.add_run(f"[{placeholder}]")
    r.italic = True
    r.font.color.rgb = GREY


def add_figure(doc: docx.Document, jinja_var: str, figure_number: int, caption: str):
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run("{{ " + jinja_var + " }}")

    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap_p.add_run(f"Figura {figure_number}. {caption}")
    cap_run.italic = True
    cap_run.font.size = Pt(9)
    cap_run.font.color.rgb = GREY


# ── main ──────────────────────────────────────────────────────────────────────


def main() -> None:
    doc = docx.Document()

    # ── title page ──────────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("RELATÓRIO DE INTELIGÊNCIA TERRITORIAL")
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = TEAL

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_p.add_run("Análise dos Eventos Relacionados à Caça")
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = ORANGE

    source_p = doc.add_paragraph()
    source_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source_p.add_run("Base de dados EarthRanger").bold = True
    period_p = doc.add_paragraph()
    period_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    period_p.add_run("Período: {{ period }}").bold = True

    doc.add_paragraph(
        "Documento temático elaborado a partir da análise integrada dos eventos relacionados à caça "
        "registrados na plataforma EarthRanger. O relatório busca identificar padrões temporais, "
        "espaciais e operacionais capazes de subsidiar o planejamento das ações de proteção ambiental."
    )
    doc.add_page_break()

    # ── 1-2: narrative-only sections ───────────────────────────────────────
    add_heading(doc, "1", "Objetivo")
    add_blank_editable_paragraph(doc, "Texto a ser preenchido pelo usuário")

    add_heading(doc, "2", "Fonte dos dados e metodologia")
    add_blank_editable_paragraph(doc, "Texto a ser preenchido pelo usuário")

    # ── 3: Caracterização geral dos registros ──────────────────────────────
    add_heading(doc, "3", "Caracterização geral dos registros")
    stats_table = doc.add_table(rows=7, cols=2)
    stats_table.style = "Light Grid Accent 1"
    stats_rows = [
        ("Eventos analisados", "{{ total_events }}"),
        ("Período", "{{ period }}"),
        ("Locais ativos", "{{ locais_ativos }}"),
        ("Locais inativos", "{{ locais_inativos }}"),
        ("Fotografias disponibilizadas", "{{ fotografias_disponibilizadas }}"),
        ("Eventos com infrator presente", "{{ infrator_presente }}"),
        ("Agrupamentos espaciais exploratórios", "{{ agrupamentos }}"),
    ]
    for i, (label, expr) in enumerate(stats_rows):
        set_cell_text(stats_table.rows[i].cells[0], label, bold=True)
        set_cell_text(stats_table.rows[i].cells[1], expr)

    doc.add_paragraph()
    effort_heading = doc.add_paragraph()
    effort_run = effort_heading.add_run("Esforço de patrulhamento e Taxa de Detecção de Estruturas Ativas (TDEA)")
    effort_run.bold = True
    doc.add_paragraph("Número de patrulhas: {{ total_patrols }}    |    Horas trabalhadas: {{ total_hours }} h")
    doc.add_paragraph(
        "TDEA = nº de estruturas ativas / km patrulhados × 100 — expressa como estruturas ativas "
        "detectadas a cada 100 km de patrulhamento."
    )
    tdea_table = doc.add_table(rows=2, cols=4)
    tdea_table.style = "Light Grid Accent 1"
    for i, h in enumerate(["Período", "Km patrulhados", "Estruturas ativas", "TDEA"]):
        set_cell_text(tdea_table.rows[0].cells[i], h, bold=True)
    for i, expr in enumerate(["{{ period }}", "{{ total_distance_km }}", "{{ estruturas_ativas }}", "{{ tdea }}"]):
        set_cell_text(tdea_table.rows[1].cells[i], expr)
    note_p = doc.add_paragraph()
    note_run = note_p.add_run(
        "Nota: adicione manualmente uma nova linha a esta tabela a cada novo período, para permitir a "
        "comparação da evolução do esforço de patrulhamento e da TDEA ao longo do tempo."
    )
    note_run.italic = True
    note_run.font.size = Pt(9)
    note_run.font.color.rgb = GREY

    # ── 4: Distribuição temporal ────────────────────────────────────────────
    add_heading(doc, "4", "Distribuição temporal")
    add_figure(doc, "temporal_chart", 1, "Distribuição mensal dos eventos relacionados à caça registrados no EarthRanger.")

    # ── 5: Caracterização operacional ───────────────────────────────────────
    add_heading(doc, "5", "Caracterização operacional")
    add_figure(doc, "situacao_chart", 2, "Situação dos locais de caça no momento do registro.")
    add_figure(doc, "tempo_chart", 3, "Tempo estimado da atividade associada aos registros.")
    add_figure(doc, "estruturas_chart", 4, "Frequência das estruturas observadas nos eventos.")
    add_figure(doc, "acoes_chart", 5, "Ações tomadas pelas equipes nos eventos registrados.")

    # ── 6: Distribuição territorial e hotspots ──────────────────────────────
    add_heading(doc, "6", "Distribuição territorial e hotspots")
    add_figure(doc, "events_map", 6, "Mapa geral dos eventos relacionados à caça fornecidos para o relatório.")
    add_figure(doc, "density_chart", 7, "Densidade relativa e agrupamentos espaciais exploratórios calculados a partir das coordenadas.")
    add_table_with_row_loop(
        doc,
        headers=["Agrupamento", "Eventos", "Locais ativos", "Coordenada central", "Referência descritiva predominante"],
        loop_var="hotspots",
        cell_exprs=[
            "{{ item['Agrupamento'] }}",
            "{{ item['Eventos'] }}",
            "{{ item['Locais ativos'] }}",
            "{{ item['Coordenada central'] }}",
            "{{ item['Referência descritiva predominante'] }}",
        ],
    )
    caveat_p = doc.add_paragraph()
    caveat_run = caveat_p.add_run(
        'A denominação "hotspot" refere-se à concentração de detecções em proximidade espacial, não '
        "sendo uma delimitação oficial de setor operacional."
    )
    caveat_run.italic = True
    caveat_run.font.size = Pt(9)
    caveat_run.font.color.rgb = GREY

    # ── 7-10: narrative-only sections ───────────────────────────────────────
    add_heading(doc, "7", "Relações com acesso, limites e cursos d'água")
    add_blank_editable_paragraph(doc, "Texto a ser preenchido pelo usuário")

    add_heading(doc, "8", "Tendências observadas")
    add_blank_editable_paragraph(doc, "Texto a ser preenchido pelo usuário")

    add_heading(doc, "9", "Avaliação Estratégica")
    add_blank_editable_paragraph(doc, "Texto a ser preenchido pelo usuário")

    add_heading(doc, "10", "Recomendações para o planejamento das próximas ações")
    add_blank_editable_paragraph(doc, "Texto a ser preenchido pelo usuário")

    # ── 11: Evidências Fotográficas Selecionadas ────────────────────────────
    add_heading(doc, "11", "Evidências Fotográficas Selecionadas")
    doc.add_paragraph(
        "As imagens a seguir ilustram estruturas e vestígios associados aos eventos analisados, "
        "com legenda indicando a data e a hora do evento associado."
    )
    add_table_with_row_loop(
        doc,
        headers=["Informações", "Imagem"],
        loop_var="photos",
        cell_exprs=["{{ item.caption }}", "{{ item.image }}"],
    )

    DST.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(DST))
    print(f"Template written to: {DST}")


if __name__ == "__main__":
    main()
